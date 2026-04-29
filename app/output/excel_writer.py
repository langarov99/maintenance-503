import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from pathlib import Path

from ..processors.field_mapper import ProductRecord

COLUMNS = [
    ("Код на продукта",         "product_code",   20),
    ("Количество",              "quantity_num",    12),
    ("М.ед.",                   "quantity_unit",   10),
    ("Цена за брой",            "price_val",       14),
    ("Валута",                  "price_cur",        8),
    ("Обща сума",               "total_val",       14),
    ("Валута ",                 "total_cur",        8),
    ("Нов продукт",             "is_new_product",  14),
    ("Име на продукта",         "product_name",    40),
    ("EAN / Баркод",            "ean",             18),
    ("Килограми (бруто/нето)",  "weight_kg",       22),
    ("Брой/части в комплект",   "parts_in_set",    20),
    ("Цвят",                    "color",           16),
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
ALT_FILL   = PatternFill("solid", fgColor="D6E4F0")
BORDER_SIDE = Side(style="thin", color="BFBFBF")
CELL_BORDER = Border(left=BORDER_SIDE, right=BORDER_SIDE,
                     top=BORDER_SIDE, bottom=BORDER_SIDE)


def _split(raw: str):
    """Split 'VALUE UNIT' → ('VALUE', 'UNIT').  Returns ('', '') if blank."""
    if not raw or not raw.strip():
        return "", ""
    m = re.match(r'^([\d.,]+)\s*([A-Za-z%]*)', raw.strip())
    if m:
        return m.group(1), m.group(2).upper()
    return raw.strip(), ""


def _get_value(rec: ProductRecord, field_name: str) -> str:
    if field_name == "quantity_num":   return _split(rec.quantity    or "")[0]
    if field_name == "quantity_unit":  return _split(rec.quantity    or "")[1]
    if field_name == "price_val":      return _split(rec.price       or "")[0]
    if field_name == "price_cur":      return _split(rec.price       or "")[1]
    if field_name == "total_val":      return _split(rec.total_price or "")[0]
    if field_name == "total_cur":      return _split(rec.total_price or "")[1]
    if field_name == "is_new_product": return "Да" if rec.is_new_product else "Не"
    return getattr(rec, field_name, None) or ""


def _apply_header_style(cell):
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = CELL_BORDER


def write_excel(records: list[ProductRecord], output_dir: str, source_filename: str,
                include_fields: set[str] | None = None) -> str:
    # Filter active columns
    active_cols = [(h, f, w) for h, f, w in COLUMNS
                   if include_fields is None or f in include_fields]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Извлечени данни"

    # Header row
    for col_idx, (header, _, width) in enumerate(active_cols, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        _apply_header_style(cell)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 30

    # Data rows
    for row_idx, rec in enumerate(records, start=2):
        fill = ALT_FILL if row_idx % 2 == 0 else PatternFill()
        for col_idx, (_, field_name, _) in enumerate(active_cols, start=1):
            value = _get_value(rec, field_name)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.border = CELL_BORDER
            if fill.fill_type:
                cell.fill = fill

    # Metadata sheet
    ws_meta = wb.create_sheet("Информация")
    ws_meta["A1"] = "Изходен файл"
    ws_meta["B1"] = source_filename
    ws_meta["A2"] = "Дата на обработка"
    ws_meta["B2"] = datetime.now().strftime("%Y-%m-%d %H:%M")
    ws_meta["A3"] = "Брой записи"
    ws_meta["B3"] = len(records)
    ws_meta["A4"] = "Метод на извличане"
    methods = list({r.extraction_method for r in records})
    ws_meta["B4"] = ", ".join(methods)

    # Freeze header
    ws.freeze_panes = "A2"

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    stem = Path(source_filename).stem
    out_path = Path(output_dir) / f"{stem}_extracted_{ts}.xlsx"
    wb.save(str(out_path))
    return str(out_path)
